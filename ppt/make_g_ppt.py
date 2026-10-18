# -*- coding: utf-8 -*-
"""生成《G计划 · 智能桌面宠物系统》比赛答辩 PPT（10 页，16:9）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 常量 ----------
SW = Inches(13.333)
SH = Inches(7.5)

INK   = RGBColor(0x0B, 0x1F, 0x2A)   # 主文字（深蓝黑）
INK2  = RGBColor(0x3A, 0x4A, 0x55)   # 次级文字
CYAN  = RGBColor(0x00, 0xB9, 0xD4)   # 主色（冰蓝）
CYAN_D = RGBColor(0x00, 0x8E, 0xA8)  # 深冰蓝
LIGHT = RGBColor(0xE6, 0xF7, 0xFA)   # 浅蓝底
GRAY  = RGBColor(0x7A, 0x88, 0x93)   # 灰
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xE8, 0x7A, 0x1E)  # 女仆暖橙
RED   = RGBColor(0xDD, 0x35, 0x4A)   # 守卫红

FONT = '微软雅黑'
TOTAL = 10

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(OUT_DIR, 'G计划_智能桌面宠物系统.pptx')

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


# ---------- 工具 ----------
def set_run(run, text, size=18, bold=False, color=INK, name=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def outline(shape, color, w=1.0):
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(w)
    return shape


def add_footer(slide, page):
    tb, tf = add_textbox(slide, Inches(0.55), Inches(7.02), Inches(5.5), Inches(0.35))
    r = tf.paragraphs[0].add_run()
    set_run(r, 'G计划 · 智能桌面宠物系统', size=10, color=GRAY)
    tb2, tf2 = add_textbox(slide, Inches(11.7), Inches(7.02), Inches(1.1), Inches(0.35))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    set_run(r2, '%02d / %02d' % (page, TOTAL), size=10, color=GRAY)


def add_title(slide, text, page):
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.52), Inches(0.1), Inches(0.62)), CYAN)
    tb, tf = add_textbox(slide, Inches(0.85), Inches(0.44), Inches(10.8), Inches(0.8))
    r = tf.paragraphs[0].add_run()
    set_run(r, text, size=28, bold=True, color=INK)
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.87), Inches(1.18), Inches(2.3), Inches(0.03)), CYAN)
    # 右上角装饰
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.66), Inches(0.45), Inches(0.42), Inches(0.026)), CYAN)
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(13.05), Inches(0.45), Inches(0.026), Inches(0.42)), CYAN)
    add_footer(slide, page)


def add_bullets(slide, items, left, top, width, height, size=17, gap=10):
    tb, tf = add_textbox(slide, left, top, width, height)
    first = True
    for it in items:
        lead, desc = it if isinstance(it, tuple) else (it, None)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.line_spacing = 1.12
        r0 = p.add_run()
        set_run(r0, '▪ ', size=size, bold=True, color=CYAN)
        r1 = p.add_run()
        if desc:
            set_run(r1, lead + '：', size=size, bold=True, color=INK)
            r2 = p.add_run()
            set_run(r2, desc, size=size, bold=False, color=INK2)
        else:
            set_run(r1, lead, size=size, bold=False, color=INK)
    return tb


def add_card(slide, left, top, width, height, title, desc, accent=CYAN):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.fill.background()
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.09), height), accent)
    tb, tf = add_textbox(slide, left + Inches(0.35), top + Inches(0.18), width - Inches(0.6), Inches(0.5))
    r = tf.paragraphs[0].add_run()
    set_run(r, title, size=18, bold=True, color=INK)
    tb2, tf2 = add_textbox(slide, left + Inches(0.35), top + Inches(0.72), width - Inches(0.6), height - Inches(0.9))
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    set_run(r2, desc, size=14, color=INK2)
    p2.line_spacing = 1.15
    return box


# ---------- 页面 ----------
def cover():
    slide = prs.slides.add_slide(BLANK)
    # 右上装饰圆环
    rings = [(5.6, 1.4), (4.2, 1.0), (2.9, 0.7)]
    for d, w in rings:
        outline(slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.8 - d / 2), Inches(0.5 - d / 2), Inches(d), Inches(d)), CYAN, w)
    # 左下淡色大环
    outline(slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2.2), Inches(5.6), Inches(5.4), Inches(5.4)), LIGHT, 1.0)

    # 顶部英文标签
    tb, tf = add_textbox(slide, Inches(0.9), Inches(0.8), Inches(8), Inches(0.4))
    r = tf.paragraphs[0].add_run()
    set_run(r, 'DESKTOP COMPANION SYSTEM', size=13, bold=True, color=CYAN)

    # 主标题
    tb, tf = add_textbox(slide, Inches(0.9), Inches(2.55), Inches(11.5), Inches(1.6))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, 'G计划 · 智能桌面宠物系统', size=46, bold=True, color=INK)

    # 分隔线
    solid(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(4.15), Inches(2.6), Inches(0.04)), CYAN)

    # 副标题
    tb, tf = add_textbox(slide, Inches(0.9), Inches(4.45), Inches(11.5), Inches(0.9))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '基于 Electron 的桌面智能伴侣与系统状态可视化', size=20, color=INK2)

    # 底部信息
    tb, tf = add_textbox(slide, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.5))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '项目答辩', size=14, bold=True, color=INK)
    r2 = p.add_run()
    set_run(r2, '　　团队 / 作者：____________', size=14, color=GRAY)


def s02():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '项目背景与意义', 2)
    items = [
        ('桌面场景高频', '桌面是日常学习与工作的主要入口，但传统桌面组件缺乏情感陪伴与主动信息呈现能力'),
        ('传统挂件局限', '多数桌面挂件功能单一、静态，难以提供拟人化、动态、可交互的体验'),
        ('状态信息割裂', 'CPU、内存、电池、网络等系统状态通常隐藏在任务栏或设置中，缺少直观、实时的可视化'),
        ('项目意义', '将「情感陪伴」与「系统状态可视化」整合于单一桌面组件，提升桌面交互体验与信息获取效率'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0), size=18, gap=16)


def s03():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '项目概述', 3)
    tb, tf = add_textbox(slide, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.7))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '基于 Electron 的跨平台桌面智能宠物系统，以拟人化角色为载体，融合陪伴交互与实时系统监控。',
            size=17, color=INK2)
    # 三角色卡片
    roles = [
        ('舰长 · 指挥官', '科技 / 指挥风格，冰蓝主题，负责数据与状态播报', CYAN),
        ('女仆 · 陪伴者', '温暖 / 欢迎风格，暖橙主题，负责问候与日常陪伴', ORANGE),
        ('守卫 · 警戒者', '安全 / 警戒风格，红色主题，负责异常状态告警', RED),
    ]
    x = Inches(0.9)
    w = Inches(3.72)
    for name, desc, c in roles:
        add_card(slide, x, Inches(2.6), w, Inches(1.6), name, desc, accent=c)
        x += w + Inches(0.17)
    tb, tf = add_textbox(slide, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.7))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '核心能力：拟人化角色切换 · 语音气泡交互 · 实时系统状态监控与告警', size=16, bold=True, color=INK)


def s04():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '系统架构', 4)
    layers = [
        ('渲染层  renderer.js + HTML/CSS', '负责角色渲染、HUD 数据展示、动效与交互逻辑'),
        ('预加载层  preload.js', '通过 contextBridge 安全暴露 API，桥接渲染层与主进程，避免直接暴露 Node 能力'),
        ('主进程  main.js', '窗口管理、系统状态采集（CPU / 内存 / 电池 / 网络）、IPC 通信'),
    ]
    y = Inches(1.6)
    h = Inches(1.15)
    for i, (t, d) in enumerate(layers):
        add_card(slide, Inches(1.1), y, Inches(11.1), h, t, d)
        if i < len(layers) - 1:
            ar = solid(slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), y + h + Inches(0.03), Inches(0.35), Inches(0.3)), GRAY)
        y += h + Inches(0.4)
    tb, tf = add_textbox(slide, Inches(1.1), Inches(6.35), Inches(11.1), Inches(0.5))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '通信机制：ipcRenderer / ipcMain 的 invoke 调用与事件推送，实现主进程到渲染层状态的实时同步。',
            size=14, color=INK2)


def s05():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '核心功能设计', 5)
    items = [
        ('角色切换', '通过快捷键在舰长 / 女仆 / 守卫三个角色间切换，主题色、立绘、台词与 HUD 联动更新'),
        ('语音气泡', '根据角色与事件（问候、低电量、断网等）动态生成差异化台词'),
        ('系统监控', '实时采集并展示 CPU 使用率、内存占用、电池电量、网络状态'),
        ('状态告警', '低电量或断网时触发气泡提醒，并联动守卫角色的警戒动效'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0), size=18, gap=16)


def s06():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '视觉与交互设计', 6)
    items = [
        ('角色独立 HUD 环', '每个角色配备一套独立 SVG 图形（科技环 / 花瓣 / 装甲），配色与边界清晰差异化'),
        ('多层动态背景', '网格、光晕、四角刻度、粒子氛围四层背景，随快捷盘开关平滑淡入淡出'),
        ('鼠标跟随 3D 微动效', '立绘根据鼠标位置做 ±2° 的伪 3D 倾斜，营造「注视鼠标」的视线跟随感'),
        ('动效体系', '呼吸、上下浮动、说话脉冲、角色专属动画通过独立属性叠加，互不干扰'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0), size=18, gap=16)


def s07():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '关键技术难点与创新点', 7)
    items = [
        ('CSS 独立变换属性解耦', '使用 translate / rotate / scale 独立属性替代单一 transform，位移、缩放、旋转互不覆盖，解决动画叠加冲突'),
        ('多层动画栈', '单元素以逗号分隔叠加多段动画，并通过 CSS 变量按角色注入专属动画，避免样式膨胀'),
        ('角色主题化', '以 CSS 自定义属性（主色、发光色、强调色）实现「一套结构、三套皮肤」的可扩展主题'),
        ('实时数据推送', '主进程定时采集系统状态并通过 IPC 事件推送，渲染层低耦合更新，兼用轮询兜底'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0), size=17, gap=16)


def s08():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '实现方案与技术选型', 8)
    items = [
        ('技术栈', 'Electron + 原生 HTML / CSS / JavaScript，无重型前端框架，降低运行开销'),
        ('窗口方案', '透明无边框窗口，支持跨平台拖拽，融入桌面环境'),
        ('通信安全', 'contextBridge + preload 隔离，渲染层仅通过受控 API 与主进程交互'),
        ('性能处理', 'requestAnimationFrame 驱动动画、will-change 提升合成层，兼顾流畅度与资源占用'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.0), size=18, gap=16)


def s09():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '成果与完成情况', 9)
    rows = [
        ('功能模块', '完成情况'),
        ('三角色系统与角色切换', '已完成'),
        ('语音气泡与角色台词', '已完成'),
        ('实时系统状态监控（CPU / 内存 / 电池 / 网络）', '已完成'),
        ('低电量 / 断网状态告警', '已完成'),
        ('角色独立 HUD 环（三套 SVG）', '已完成'),
        ('多层动态背景（网格 / 光晕 / 刻度 / 粒子）', '已完成'),
        ('鼠标跟随 3D 微动效与呼吸浮动', '已完成'),
    ]
    left, top, width = Inches(1.1), Inches(1.6), Inches(11.1)
    n_rows, n_cols = len(rows), 2
    gframe = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(4.9))
    table = gframe.table
    table.columns[0].width = Inches(7.5)
    table.columns[1].width = Inches(3.6)
    for i in range(n_rows):
        table.rows[i].height = Inches(0.6)
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            text = rows[r][c]
            if r == 0:
                set_cell(cell, text, size=16, bold=True, color=WHITE, align=PP_ALIGN.LEFT, fill=CYAN_D)
            else:
                fill = LIGHT if r % 2 == 0 else WHITE
                if c == 1:
                    set_cell(cell, text, size=15, bold=True, color=CYAN_D, align=PP_ALIGN.CENTER, fill=fill)
                else:
                    set_cell(cell, text, size=15, bold=False, color=INK, align=PP_ALIGN.LEFT, fill=fill)


def s10():
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, '总结与展望', 10)
    tb, tf = add_textbox(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(2.0))
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, '总结', size=18, bold=True, color=INK)
    tb2, tf2 = add_textbox(slide, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.2))
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    set_run(r2, '本项目实现了一款集陪伴、可视化、告警于一体的桌面宠物系统，工程结构清晰、动效解耦、主题可扩展。',
            size=16, color=INK2)
    p2.line_spacing = 1.2
    tb3, tf3 = add_textbox(slide, Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.5))
    p3 = tf3.paragraphs[0]
    r3 = p3.add_run()
    set_run(r3, '展望', size=18, bold=True, color=INK)
    items = [
        ('接入 AI 对话能力', '结合大语言模型，实现更自然的拟人化对话'),
        ('扩展角色与皮肤', '在主题化框架下低成本扩展更多角色与视觉皮肤'),
        ('更丰富的状态可视化', '扩展网络、磁盘、进程等更多维度的可视化呈现'),
        ('跨平台打包分发', '完成 Windows / macOS / Linux 打包，降低使用门槛'),
    ]
    add_bullets(slide, items, Inches(0.9), Inches(4.05), Inches(11.5), Inches(2.6), size=16, gap=8)


def set_cell(cell, text, size=15, bold=False, color=INK, align=PP_ALIGN.LEFT, fill=None):
    cell.margin_left = Inches(0.14)
    cell.margin_right = Inches(0.14)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size=size, bold=bold, color=color)


# ---------- 组装 ----------
cover()
s02()
s03()
s04()
s05()
s06()
s07()
s08()
s09()
s10()

os.makedirs(OUT_DIR, exist_ok=True)
prs.save(OUT_PATH)
print('OK ->', OUT_PATH)
